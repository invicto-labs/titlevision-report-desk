"""Persist reviewed attribution by native error ID; never import source decisions."""
from zipfile import ZipFile
from io import BytesIO
import copy, json, re
import openpyxl
import report_layout

EDITABLE = (3, 18, 20, 22)
TEAMS = ('', *report_layout.TEAMS)
ID_HEADER = 'TitleVision Error ID'

def initialize(c):
    c.execute('CREATE TABLE IF NOT EXISTS main_edits (month TEXT, source_id TEXT, fields TEXT NOT NULL, updated TEXT NOT NULL, PRIMARY KEY(month,source_id))')

def contributor(row):
    team = report_layout.normalize_team((row[3] or '').strip())
    return team if team in ('Triage','VM Team','Leadership') else (row[18] or '') if team == 'Search' else (row[20] or '') if team in ('Type','Typing') else ''

def load(db, month):
    with db() as c:
        return {r['source_id']: json.loads(r['fields']) for r in c.execute('SELECT * FROM main_edits WHERE month=?',(month,))}

def apply(rows, audit, edits):
    if len(rows) != len(audit): raise ValueError('Contributor corrections have no matching error identity.')
    rows = copy.deepcopy(rows)
    overrides = []
    for row, item in zip(rows, audit):
        saved = edits.get(item['sourceId'], {})
        for field in (3,18,20):
            key = str(field)
            if field == 3 and key not in saved: key = '22'  # older saved Team edits
            if key in saved: row[field] = saved[key]
        row[3] = report_layout.normalize_team(row[3])
        override = report_layout.normalize_override(saved.get('final'))
        if override and override['team'] != row[3]: override = None
        row[22] = override['name'] if override else contributor(row)
        overrides.append(override)
    return rows, overrides

def read_upload(content):
    if not content or len(content) > 12*1024*1024: raise ValueError('Choose an Excel workbook smaller than 12 MB.')
    try:
        with ZipFile(BytesIO(content)) as z:
            if len(z.infolist()) > 1500 or sum(i.file_size for i in z.infolist()) > 80*1024*1024:
                raise ValueError('The Excel workbook is too large to read safely.')
        return openpyxl.load_workbook(BytesIO(content), data_only=False, read_only=True, keep_links=False)
    except Exception as error:
        raise ValueError('Cannot read this Excel workbook: '+str(error)) from error

def save(root, data, db, run_process, month, content):
    import monthly_sync as sync
    import main_workbook as main
    sync.guard(month)
    wb = read_upload(content)
    try:
        if '_ReportDesk' not in wb or 'SP 2' not in wb:
            raise ValueError('Download the latest Main workbook with this app version before editing it. Older copies have no error IDs.')
        meta = wb['_ReportDesk']
        revision = meta['B2'].value
        if meta['B1'].value != month or not isinstance(revision,str) or not re.fullmatch('[0-9a-f]{32}',revision):
            raise ValueError('The uploaded workbook belongs to another month or has invalid metadata.')
        base_folder = data/'main'/month/revision
        if not (base_folder/'validation.json').exists() or json.loads((base_folder/'validation.json').read_text(encoding='utf8')).get('passed') is not True:
            raise ValueError('This workbook was not generated and verified by this installation.')
        base = json.loads((base_folder/'payload.json').read_text(encoding='utf8'))
        with db() as c: current_file = main.download(data,c,month)
        current = json.loads((current_file.parent/'payload.json').read_text(encoding='utf8'))
        if not base.get('editIds') or not current.get('editIds'): raise ValueError('Refresh this Main workbook, download it again, and then edit it.')
        baseline = dict(zip(base['editIds'],report_layout.current_rows(base['data'][0],base['data'][1:])))
        latest = dict(zip(current['editIds'],report_layout.current_rows(current['data'][0],current['data'][1:])))
        sheet = wb['SP 2']
        if sheet.max_row > len(baseline)+1 or sheet.max_column != len(base['data'][0])+1:
            raise ValueError('Do not add or remove rows or columns. Only edit the attribution fields.')
        iterator = sheet.iter_rows(values_only=False)
        headers = [cell.value for cell in next(iterator)]
        if headers != [*base['data'][0],ID_HEADER]: raise ValueError('Workbook columns changed. Keep the original headers.')
        edits = load(db,month); changes = 0; seen = set(); detail = []
        for cells in iterator:
            cells = report_layout.current_row(cells[:-1]) + [cells[-1]]
            key = cells[27].value
            if not isinstance(key,str) or key not in baseline or key not in latest or key in seen:
                raise ValueError('An error ID is missing, repeated or no longer approved. Download the latest workbook.')
            seen.add(key); original = baseline[key]; now = latest[key]; values = [c.value for c in cells[:27]]
            # Prove that sorting kept the ID on its original source row.
            for j in (1,2,4,5,6,7,8,9,11,12,16,17):
                if (values[j] or '') != (original[j] or ''):
                    raise ValueError('Source columns were edited or an error ID moved to another row. Only contributor fields can be saved.')
            for j in (10,15):
                actual = values[j]
                if hasattr(actual,'isoformat'):
                    from datetime import datetime
                    actual = (actual-datetime(1899,12,30)).total_seconds()/86400
                if actual != original[j] and not (isinstance(actual,(float,int)) and isinstance(original[j],(float,int)) and abs(actual-original[j])<1e-7):
                    raise ValueError('Created Date or Error Committed Date was changed. Keep the source dates.')
            saved = dict(edits.get(key,{})); changed = []
            for j in (3,18,20):
                value = values[j] or ''
                if cells[j].data_type in ('f','e') or not isinstance(value,str) or len(value)>255:
                    raise ValueError('Contributor fields must be plain text, up to 255 characters.')
                if j==3:
                    value = report_layout.normalize_team(value)
                    values[3] = value
                    if value not in TEAMS: raise ValueError('Choose Search, Typing, Triage, VM Team or Leadership.')
                if value != (original[j] or ''):
                    if (now[j] or '') not in ((original[j] or ''),value): raise ValueError('A newer saved correction conflicts with this copy. Download the latest workbook and edit it again.')
                    saved[str(j)] = value; changed.append(j)
            # A formula follows Team/Searcher/Typer. A typed final name is an explicit override.
            if cells[22].data_type != 'f':
                value = values[22] or ''
                if cells[22].data_type == 'e' or not isinstance(value,str) or len(value)>255: raise ValueError('Final Error Contributor must be plain text.')
                if value == 'VM team': value = 'VM Team'
                if value != (original[22] or ''):
                    if (now[22] or '') not in ((original[22] or ''),value): raise ValueError('A newer contributor correction conflicts with this copy. Download the latest workbook.')
                    saved['final'] = {'team':values[3] or '', 'name':value};changed.append(22)
            if changed:
                edits[key] = saved;changes += len(changed);detail.append({'sourceId':key,'columns':changed})
        if seen != set(baseline): raise ValueError('Rows were removed. Upload the complete workbook, including filtered-out rows.')
    finally: wb.close()
    if not changes: return {'message':'No new contributor edits found. Your Main workbook is unchanged.'}
    rows, overrides = apply(report_layout.current_rows(current['data'][0],current['data'][1:]), [{'sourceId':i} for i in current['editIds']], edits)
    collection = json.loads((current_file.parent/'collection.json').read_text(encoding='utf8'))
    days = sorted(d for d in sync.selected_days(db) if d.startswith(month+'-'))
    audit = [{'sourceId':i,'manualEdits':edits.get(i,{})} for i in current['editIds']]
    book = sync.prepare(root,data,run_process,month,days,rows,audit,collection.get('syncThrough'),collection.get('syncedAt'),overrides)
    (data/book[1]/'manual-edit-audit.json').write_text(json.dumps({'baseRevision':revision,'changes':detail}),encoding='utf8')
    # Save corrections and published workbook together only after all checks pass.
    with db() as c:
        c.execute('INSERT OR REPLACE INTO main_books VALUES (?,?,?,?,?,?,?,?)',book)
        c.executemany('INSERT OR REPLACE INTO main_edits VALUES (?,?,?,?)',[(month,k,json.dumps(v),sync.stamp()) for k,v in edits.items()])
    return {'message':f'Saved {changes} contributor field change(s). Future monthly updates will keep these corrections.'}
