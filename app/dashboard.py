"""Read the saved monthly Excel sheet for the dashboard; never modify or sync it."""
from collections import Counter
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from datetime import date,datetime
from decimal import Decimal,InvalidOperation
import copy,re,threading,math
import openpyxl
import main_workbook,report_layout,shared_workbook

TEAMS=('Search','Typing','VM Team','Leadership','Triage')
_cache={}
_lock=threading.Lock()


def error_day(value,epoch):
    if isinstance(value,datetime):return value.date().isoformat()
    if isinstance(value,date):return value.isoformat()
    if isinstance(value,(int,float)) and not isinstance(value,bool) and math.isfinite(value):
        try:return openpyxl.utils.datetime.from_excel(value,epoch).date().isoformat()
        except (ValueError,OverflowError,AttributeError):pass
    if isinstance(value,str):
        try:return datetime.fromisoformat(value.strip()).date().isoformat()
        except ValueError:pass
    raise ValueError('An SP 2 error has an invalid Created Date. Review the workbook.')


def error_points(value):
    if isinstance(value,bool) or not isinstance(value,(int,float,Decimal)):
        raise ValueError('An SP 2 error has invalid Points. Review the workbook.')
    try:points=Decimal(str(value))
    except InvalidOperation:raise ValueError('An SP 2 error has invalid Points. Review the workbook.')
    if not points.is_finite() or points<0:raise ValueError('An SP 2 error has invalid Points. Review the workbook.')
    return points


def totals(counts,points):
    teams=[*TEAMS,*[extra for extra in ('Unassigned','Other') if counts[extra]]]
    return {'rows':[{'team':team,'count':counts[team],'points':float(points[team])} for team in teams],
            'total':sum(counts.values()),'points':float(sum(points.values(),Decimal(0)))}


def workbook_path(data,db,month):
    if not isinstance(month,str) or not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])',month):
        raise ValueError('Choose a valid workbook month.')
    with db() as c:local=main_workbook.download(data,c,month)
    if shared_workbook.configured(db):
        return shared_workbook.target(db,month),'shared'
    return local,'local'


def signature(path):
    stat=path.stat()
    return stat.st_mtime_ns,stat.st_size,stat.st_ctime_ns


def read_file(path):
    before=signature(path)
    if before[1]>12*1024*1024:raise ValueError('Workbook is too large for the dashboard.')
    content=path.read_bytes()
    if signature(path)!=before:raise ValueError('Workbook is being saved. The dashboard will retry shortly.')
    return content,before


def parse(content):
    with ZipFile(BytesIO(content)) as z:
        if len(z.infolist())>1500 or sum(i.file_size for i in z.infolist())>80*1024*1024:
            raise ValueError('Workbook is too large for the dashboard.')
    wb=openpyxl.load_workbook(BytesIO(content),read_only=True,data_only=False,keep_links=False)
    try:
        if 'SP 2' not in wb:raise ValueError('The workbook has no SP 2 sheet.')
        sheet=wb['SP 2']
        if sheet.max_row>50000:raise ValueError('Workbook has too many rows for the dashboard.')
        iterator=sheet.iter_rows(values_only=True);headers=list(next(iterator,()))
        if headers and headers[-1]=='TitleVision Error ID':headers=headers[:-1]
        if headers not in (report_layout.HEADERS,report_layout.PREVIOUS_HEADERS,report_layout.LEGACY_HEADERS):
            raise ValueError('Workbook columns changed. Keep the original SP 2 headers.')
        team_col=headers.index('Team');points_col=headers.index('Points');date_col=headers.index('Created Date')
        counts=Counter();points=Counter();other=Counter();daily={}
        for row in iterator:
            # The empty workbook has a preformatted formula row. It is not an error.
            if row[1] in (None,''):
                if any(row[j] not in (None,'') for j in (0,2,5,6,8,10)):
                    raise ValueError('An SP 2 error row has no order number. Review the workbook.')
                continue
            team=row[team_col]
            if team is not None and (not isinstance(team,str) or team.startswith('=')):
                raise ValueError('Choose a Team from the Excel dropdown; formulas are not supported there.')
            team=report_layout.normalize_team((team or '').strip())
            group=team if team in TEAMS else 'Unassigned' if not team else 'Other'
            amount=error_points(row[points_col]);day=error_day(row[date_col],wb.epoch)
            counts[group]+=1;points[group]+=amount
            day_counts,day_points=daily.setdefault(day,(Counter(),Counter()))
            day_counts[group]+=1;day_points[group]+=amount
            if group=='Other':other[team]+=1
        latest=max(daily) if daily else None
        return {**totals(counts,points),'otherTeams':dict(other),
                'latest':{'date':latest,**totals(*(daily[latest] if latest else (Counter(),Counter())))}}
    finally:wb.close()


def summary(data,db,month):
    try:
        path,source=workbook_path(data,db,month)
        sig=signature(path);key=str(path)
        with _lock:cached=_cache.get(key)
        if cached and cached[0]==sig:result=copy.deepcopy(cached[1])
        else:
            content,sig=read_file(path);result=parse(content)
            with _lock:
                if len(_cache)>36:_cache.clear()
                _cache[key]=(sig,copy.deepcopy(result))
        return {**result,'available':True,'month':month,'source':source,'fileName':path.name,'savedAt':sig[0]/1e9}
    except Exception as error:
        # Never show old cached totals or silently fall back to a local copy when the share is unavailable.
        message=str(error) if isinstance(error,ValueError) else 'Cannot read the saved workbook. Check the shared folder and try again.'
        return {'available':False,'month':month,'message':message}
