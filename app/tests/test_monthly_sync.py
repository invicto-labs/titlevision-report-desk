from contextlib import contextmanager
from datetime import date,datetime
from pathlib import Path
from unittest.mock import patch
import copy,hashlib,json,shutil,sqlite3,subprocess,sys,unittest,uuid,zipfile
from xml.etree import ElementTree as ET
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import main_workbook as main
import monthly_sync as sync
import manual_edits
from status_history import fields as status_fields
ROOT=Path(__file__).resolve().parents[1]
SITE_STATUSES=('New','Accepted','Auto-Accepted','Disputed','Non-Chargeable','Chargeable')

class MonthlySyncTests(unittest.TestCase):
 def setUp(self):
  self.data=ROOT/'tests/results'/('phase2-'+uuid.uuid4().hex);self.data.mkdir(parents=True)
  with self.db() as c:
   c.execute('CREATE TABLE runs (id TEXT PRIMARY KEY,start TEXT,end TEXT,status TEXT,message TEXT,created TEXT,count INTEGER,points REAL,issues INTEGER,scheduled_day TEXT)');main.initialize(c)
  self.clock=patch.object(sync,'today',return_value=date(2026,10,4));self.clock.start();self.addCleanup(self.clock.stop)
  self.live=[];self.fetches=[]
 @contextmanager
 def db(self):
  c=sqlite3.connect(self.data/'test.sqlite');c.row_factory=sqlite3.Row
  try:
   with c:yield c
  finally:c.close()
 def source(self,id,day,points=3,status='New'):
  stamp=datetime.fromisoformat(day).strftime('%m/%d/%Y')
  v=['']*20
  for i,value in {0:status,1:'SAME-ORDER',6:'Searching',7:'Error example',9:'Full Title',12:stamp+' 12:00:00 PM',15:str(points),16:stamp,17:stamp,18:'EPON-1'}.items():v[i]=value
  return {'sourceId':str(id),'url':'https://tv.datatracetitle.com/OrderOverview.aspx?PublicOrderId=example-order','values':v}
 def seed(self,start,end,source):
  rid=uuid.uuid4().hex;folder=self.data/'runs'/rid;folder.mkdir(parents=True)
  rows=[]
  for s in source:
   v=s['values'];r=['']*28;r[0]=v[0];r[1]=v[1];r[6]=v[6];r[7]=v[7];r[9]=v[9]
   r[10]=(datetime.strptime(v[12],'%m/%d/%Y %I:%M:%S %p')-datetime(1899,12,30)).total_seconds()/86400
   r[13]=float(v[15]);r[14]=r[15]=int(r[10]);r[16]=v[18];r[18]='Searcher';r[20]='Typer';r[22]='Search';r[23]='Searcher';rows.append(r)
  collection={'start':start,'end':end,'count':len(rows),'points':sum(r[13] for r in rows)}
  for name,obj in [('source.json',source),('payload.json',main.make_payload(ROOT,rows)),('collection.json',collection),('validation.json',{'passed':True,**collection})]:(folder/name).write_text(json.dumps(obj),encoding='utf8')
  with self.db() as c:c.execute('INSERT INTO runs VALUES (?,?,?,?,?,?,?,?,?,?)',(rid,start,end,'complete','verified','now',len(rows),collection['points'],0,None))
  return rid
 def fetch(self,start,end,folder):
  self.fetches.append((start,end))
  source=[copy.deepcopy(record) for record in self.live if start<=datetime.strptime(record['values'][12],'%m/%d/%Y %I:%M:%S %p').date().isoformat()<=end]
  for record in source:
   status=record['values'][0]
   if status!='New' and 'statusHistory' not in record:
    key=record['sourceId'];order='example-order'
    record['statusHistory']={'schema':1,'sourceId':key,'orderId':order,'url':f'https://tv.datatracetitle.com/UserErrors.aspx?EditMode=Status&UserErrorId={key}&PublicOrderId={order}',
     'status':status,'vendor':record['values'][2],'vendorUsers':['VendorUser'],'comments':[['10/2/2026 1:00:00 AM','VendorUser','Our reply'],['10/3/2026 1:00:00 AM','ClientUser','Client reply']]}
   record['statusFields']=status_fields(record,required=True)
  (folder/'source.json').write_text(json.dumps(source),encoding='utf8');(folder/'source.xlsx').write_bytes(b'verified export fixture')
  (folder/'collection.json').write_text(json.dumps({'start':start,'end':end,'count':len(source),'points':sum(float(s['values'][15]) for s in source)}))
  receipt={'count':len(source),'sourceSha256':hashlib.sha256((folder/'source.json').read_bytes()).hexdigest(),'exportSha256':hashlib.sha256((folder/'source.xlsx').read_bytes()).hexdigest()}
  (folder/'reconciled.json').write_text(json.dumps(receipt))
 def build(self,args,folder):
  if str(args[1]).endswith('finish.py'):
   collection=json.loads((folder/'collection.json').read_text());(folder/'report.xlsx').write_bytes(b'validated fixture')
   (folder/'validation.json').write_text(json.dumps({'passed':True,'pivots':4,**collection}))
 def real_build(self,args,folder):
  result=subprocess.run([str(x) for x in args],capture_output=True,text=True)
  if result.returncode:raise ValueError(result.stdout+result.stderr)
 def add(self,rid,builder=None):return sync.add(ROOT,self.data,self.db,builder or self.build,self.fetch,rid)
 def book(self):
  with self.db() as c:path=main.download(self.data,c,'2026-10')
  return path,json.loads((path.parent/'payload.json').read_text())['data'][1:]
 def test_strict_start_and_month_boundary(self):
  with patch.object(sync,'today',return_value=date(2026,8,31)):
   with self.assertRaisesRegex(ValueError,'1 September'):sync.create(ROOT,self.data,self.db,self.build,'2026-10')
  for month in ['2026-08','2026-11','../../runs','2026-99']:
   with self.assertRaises(ValueError):sync.create(ROOT,self.data,self.db,self.build,month)
  with patch.object(sync,'today',return_value=date(2026,10,1)):
   sync.create(ROOT,self.data,self.db,self.build,'2026-10')
   with self.assertRaisesRegex(ValueError,'No completed day'):sync.refresh(ROOT,self.data,self.db,self.build,self.fetch,'2026-10')
  self.assertEqual(self.fetches,[])
 def test_september_refresh_and_friday_to_sunday_month_boundary(self):
  september=self.source(11,'2026-09-30');october=self.source(12,'2026-10-01')
  self.live=[september,october]
  run=self.seed('2026-09-30','2026-10-01',[september,october])
  sync.add(ROOT,self.data,self.db,self.build,self.fetch,run)
  with self.db() as c:
   books={b['month']:b for b in main.books(c)}
  self.assertEqual(set(books),{'2026-09','2026-10'})
  self.assertEqual((books['2026-09']['count'],books['2026-10']['count']),(1,1))
  self.assertIn(('2026-09-01','2026-09-30'),self.fetches)
  self.assertIn(('2026-10-01','2026-10-03'),self.fetches)
 def test_earlier_dispute_is_found_and_refresh_changes_client_status(self):
  old=self.source(1,'2026-09-20',status='Disputed');self.live=[copy.deepcopy(old)]
  sync.add(ROOT,self.data,self.db,self.build,self.fetch,self.seed('2026-09-20','2026-09-20',[old]))
  self.assertEqual(sync.disputed_months(self.data,self.db),['2026-09'])
  with patch.object(sync,'today',return_value=date(2026,10,4)):
   self.live[0]['values'][0]='Non-Chargeable'
   sync.refresh(ROOT,self.data,self.db,self.build,self.fetch,'2026-09')
  with self.db() as c:file=main.download(self.data,c,'2026-09')
  row=json.loads((file.parent/'payload.json').read_text())['data'][1]
  self.assertEqual(row[0],'Non-Chargeable')
  self.assertEqual(row[24:28],['Disputed','Our reply','Non-Chargeable','Client reply'])
  self.assertEqual(sync.disputed_months(self.data,self.db),[])
 def test_exact_error_ids_preserve_repeated_orders_and_actual_nonchargeable_points(self):
  first=[self.source(1,'2026-10-01'),self.source(2,'2026-10-01')]
  second=[self.source(3,'2026-10-02'),self.source(4,'2026-10-02')]
  a=self.seed('2026-10-01','2026-10-01',first);b=self.seed('2026-10-02','2026-10-02',second)
  self.live=copy.deepcopy(first+second);self.add(a)
  for s,status,points in zip(self.live,['New','Disputed','Chargeable','Non-Chargeable'],[3,2,4,.5]):s['values'][0]=status;s['values'][15]=str(points)
  self.add(b);_,rows=self.book()
  self.assertEqual(len(rows),4);self.assertEqual([r[0] for r in rows],['New','Disputed','Chargeable','Non-Chargeable'])
  self.assertEqual([r[13] for r in rows],[3,2,4,.5]);self.assertTrue(all(r[22]=='Search' and r[23]=='Searcher' for r in rows))
  self.assertEqual(self.fetches[-1],('2026-10-01','2026-10-03'))
  self.add(b);self.assertEqual(len(self.book()[1]),4)
 def test_private_roster_renames_existing_rows_without_losing_manual_corrections(self):
  original=self.source(1,'2026-10-01');original['id']='local-one'
  self.live=[copy.deepcopy(original)]
  rid=self.seed('2026-10-01','2026-10-01',[original]);folder=self.data/'runs'/rid
  (folder/'staff.json').write_text(json.dumps({'local-one':{'names':['KishoreK_ADSSearchType','DeepikaK_ADSSearchType']}}),encoding='utf8')
  self.add(rid)
  self.assertEqual(self.book()[1][0][18:24],['Searcher','','Typer','','Search','Searcher'])
  mapping={'schema':2,'search':{'kishorek':{'id':'INV060','name':'Kishore R'}},'type':{'deepikak':{'id':'INV160','name':'Kanna Deepika'}},'legacy':{}}
  (self.data/'names.json').write_text(json.dumps(mapping),encoding='utf8')
  sync.refresh(ROOT,self.data,self.db,self.build,self.fetch,'2026-10')
  row=self.book()[1][0]
  self.assertEqual((row[18],row[20],row[23]),('Kishore R','Kanna Deepika','Kishore R'))
  with self.db() as c:c.execute('INSERT INTO main_edits VALUES (?,?,?,?)',('2026-10','1',json.dumps({'18':'Reviewed Searcher'}),'now'))
  sync.refresh(ROOT,self.data,self.db,self.build,self.fetch,'2026-10')
  row=self.book()[1][0]
  self.assertEqual((row[18],row[20],row[23]),('Reviewed Searcher','Kanna Deepika','Reviewed Searcher'))
 def test_refresh_updates_approved_rows_without_adding_unapproved_rows(self):
  approved=self.source(1,'2026-10-01');pending=self.source(2,'2026-10-02')
  self.live=[approved,pending];self.add(self.seed('2026-10-01','2026-10-01',[approved]))
  rid=self.seed('2026-10-02','2026-10-02',[pending]);main.decide(ROOT,self.data,self.db,self.build,rid,False)
  self.live[0]['values'][0]='Non-Chargeable';self.live[0]['values'][15]='1';self.live[0]['values'][16]='10/03/2026'
  sync.refresh(ROOT,self.data,self.db,self.build,self.fetch,'2026-10')
  rows=self.book()[1];self.assertEqual(len(rows),1);self.assertEqual(rows[0][0],'Non-Chargeable');self.assertEqual(rows[0][13],1)
 def test_all_site_status_transitions_and_future_status_are_preserved(self):
  original=self.source(1,'2026-10-01');self.live=[copy.deepcopy(original)]
  self.add(self.seed('2026-10-01','2026-10-01',[original]))
  # Statuses are source values, not an enum: a later site addition must survive too.
  for points,status in enumerate((*SITE_STATUSES,'Future site status')):
   with self.subTest(status=status):
    self.live[0]['values'][0]=status;self.live[0]['values'][15]=str(points)
    self.live[0]['values'][16]='10/03/2026'
    sync.refresh(ROOT,self.data,self.db,self.build,self.fetch,'2026-10')
    _,rows=self.book();self.assertEqual(len(rows),1)
    self.assertEqual(rows[0][0],status);self.assertEqual(rows[0][13],points)
    self.assertEqual(rows[0][14],(datetime(2026,10,3)-datetime(1899,12,30)).days)
    self.assertEqual(rows[0][22:24],['Search','Searcher'])
 def test_decisions_comments_refresh_even_without_status_change_and_missing_history_retains_book(self):
  original=self.source(1,'2026-10-01');self.live=[copy.deepcopy(original)];rid=self.seed('2026-10-01','2026-10-01',[original]);self.add(rid)
  current=self.live[0];current['values'][0]='Non-Chargeable'
  sync.refresh(ROOT,self.data,self.db,self.build,self.fetch,'2026-10')
  self.assertEqual(self.book()[1][0][24:28],['Disputed','Our reply','Non-Chargeable','Client reply'])
  folder=self.data/'probe';folder.mkdir();self.fetch('2026-10-01','2026-10-03',folder)
  current['statusHistory']=json.loads((folder/'source.json').read_text())[0]['statusHistory']
  current['statusHistory']['comments'][-1][2]='Training, reason accepted'
  sync.refresh(ROOT,self.data,self.db,self.build,self.fetch,'2026-10')
  self.assertEqual(self.book()[1][0][27],'Training, reason accepted')
  before=self.book()[0]
  def incomplete(a,b,f):
   self.fetch(a,b,f)
   rows=json.loads((f/'source.json').read_text());rows[0].pop('statusHistory');(f/'source.json').write_text(json.dumps(rows))
   receipt=json.loads((f/'reconciled.json').read_text());receipt['sourceSha256']=hashlib.sha256((f/'source.json').read_bytes()).hexdigest();(f/'reconciled.json').write_text(json.dumps(receipt))
  with self.assertRaisesRegex(ValueError,'history is missing'):sync.refresh(ROOT,self.data,self.db,self.build,incomplete,'2026-10')
  self.assertEqual(self.book()[0],before)
 def test_delete_keeps_daily_reports_and_allows_reapproval(self):
  s=self.source(1,'2026-10-01');self.live=[s];rid=self.seed('2026-10-01','2026-10-01',[s]);self.add(rid)
  sync.delete(self.data,self.db,'2026-10')
  with self.db() as c:
   self.assertEqual(main.books(c),[]);self.assertEqual(c.execute('SELECT COUNT(*) FROM runs').fetchone()[0],1)
   self.assertIsNone(c.execute('SELECT * FROM main_choices WHERE run_id=?',(rid,)).fetchone())
  self.assertTrue((self.data/'runs'/rid/'source.json').exists())
  sync.create(ROOT,self.data,self.db,self.build,'2026-10');self.assertEqual(self.book()[1],[])
  self.add(rid);self.assertEqual(len(self.book()[1]),1)
 def test_missing_error_failed_export_or_build_retains_previous_workbook(self):
  s=self.source(1,'2026-10-01');self.live=[s];rid=self.seed('2026-10-01','2026-10-01',[s]);self.add(rid);before=self.book()[0]
  for mode in ['missing','export','build']:
   self.live=[] if mode=='missing' else [s]
   def fetch(a,b,f):
    if mode=='export':raise ValueError('Export failed')
    self.fetch(a,b,f)
   def build(args,f):raise ValueError('Build failed')
   with self.assertRaises(ValueError):sync.refresh(ROOT,self.data,self.db,build if mode=='build' else self.build,fetch,'2026-10')
   self.assertEqual(self.book()[0],before)
   with self.db() as c:self.assertEqual(main.books(c)[0]['syncStatus'],'failed')
 def test_invalid_changed_or_duplicate_identity_is_not_guessed(self):
  for mode in ['missing-id','duplicate-id','changed-order','tampered-export']:
   with self.subTest(mode=mode):
    s=self.source(1,'2026-10-01');rid=self.seed('2026-10-01','2026-10-01',[s]);self.live=[copy.deepcopy(s)]
    if mode=='missing-id':self.live[0].pop('sourceId')
    if mode=='duplicate-id':self.live.append(copy.deepcopy(s))
    if mode=='changed-order':self.live[0]['values'][1]='OTHER-ORDER'
    def fetch(a,b,f):
     self.fetch(a,b,f)
     if mode=='tampered-export':(f/'source.xlsx').write_bytes(b'tampered')
    with self.assertRaises(ValueError):sync.add(ROOT,self.data,self.db,self.build,fetch,rid)
    with self.db() as c:self.assertIsNone(c.execute('SELECT * FROM main_choices WHERE run_id=?',(rid,)).fetchone())
 def test_missing_old_report_identity_requires_recollection(self):
  s=self.source(1,'2026-10-01');s.pop('sourceId');rid=self.seed('2026-10-01','2026-10-01',[s]);self.live=[]
  with self.assertRaisesRegex(ValueError,'Collect its dates again'):self.add(rid)
  self.assertEqual(self.fetches,[])
 def test_real_pivots_contain_refreshed_status_and_points(self):
  import openpyxl
  source=[self.source(i+1,'2026-10-01') for i in range(len(SITE_STATUSES))];rid=self.seed('2026-10-01','2026-10-01',source)
  self.live=copy.deepcopy(source);points=[3,2,1,4,0,5]
  for entry,status,value in zip(self.live,SITE_STATUSES,points):entry['values'][0]=status;entry['values'][15]=str(value)
  self.add(rid,self.real_build);file,rows=self.book();wb=openpyxl.load_workbook(file,data_only=True)
  self.assertEqual(wb['Summary']['A2'].value,'October 2026 Main Workbook')
  self.assertEqual([wb['SP 2'].cell(i+2,1).value for i in range(6)],list(SITE_STATUSES))
  self.assertEqual([wb['SP 2'].cell(i+2,14).value for i in range(6)],points)
  self.assertEqual([wb['SP 2'].cell(7,j).value for j in range(25,29)],['Disputed','Our reply','Chargeable','Client reply'])
  self.assertEqual(len(wb['Summary']._pivots),4)
  for pivot in wb['Summary']._pivots:
   self.assertEqual(pivot.cache.recordCount,6)
  with zipfile.ZipFile(file) as archive:
   ns={'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
   caches=[name for name in archive.namelist() if name.startswith('xl/pivotCache/pivotCacheDefinition') and name.endswith('.xml')]
   self.assertTrue(caches)
   for name in caches:
    field=ET.fromstring(archive.read(name)).find('x:cacheFields/x:cacheField',ns)
    self.assertEqual({s.get('v') for s in field.findall('x:sharedItems/x:s',ns)},set(SITE_STATUSES))
  totals={r[4]:tuple(r[5:7]) for r in wb['Summary'].values if r[4] in SITE_STATUSES}
  self.assertEqual(totals,{status:(1,value) for status,value in zip(SITE_STATUSES,points)})
  fixture=ROOT/'tests/results'/('report-phase2-'+uuid.uuid4().hex);fixture.mkdir();shutil.copy2(file,fixture/'report.xlsx')
  print('PHASE2_FIXTURE='+str(fixture/'report.xlsx'))

 def edited_copy(self,change):
  import openpyxl
  from io import BytesIO
  file,_=self.book();wb=openpyxl.load_workbook(file)
  change(wb);out=BytesIO();wb.save(out);wb.close();return out.getvalue()

 def test_manual_edits_survive_next_day_refresh_and_reapproval(self):
  import openpyxl
  first=[self.source(1,'2026-10-01'),self.source(2,'2026-10-01')]
  self.live=copy.deepcopy(first);a=self.seed('2026-10-01','2026-10-01',first);self.add(a,self.real_build)
  def change(wb):
   sheet=wb['SP 2'];sheet['W2']='Type';sheet['U2']='Corrected Typer';sheet['X3']='Reviewed Contributor'
   # Sorting must move the hidden native ID with its row.
   values=[list(row) for row in sheet.values][1:]
   for i,row in enumerate(reversed(values),2):
    for j,value in enumerate(row,1):sheet.cell(i,j,value)
  result=manual_edits.save(ROOT,self.data,self.db,self.real_build,'2026-10',self.edited_copy(change))
  self.assertIn('Saved 3',result['message'])
  next_day=self.source(3,'2026-10-02');self.live.append(next_day)
  self.live[0]['values'][0]='Non-Chargeable';self.live[0]['values'][15]='1'
  b=self.seed('2026-10-02','2026-10-02',[next_day]);self.add(b,self.real_build)
  file,rows=self.book();self.assertEqual(len(rows),3)
  self.assertEqual(rows[0][20:24],['Corrected Typer','','Type','Corrected Typer'])
  self.assertEqual(rows[1][23],'Reviewed Contributor');self.assertEqual(rows[2][23],'Searcher')
  self.assertEqual(rows[0][24:28],['Disputed','Our reply','Non-Chargeable','Client reply'])
  # Recollect the edited error's dates: corrections follow the error ID.
  self.add(self.seed('2026-10-01','2026-10-01',first),self.real_build)
  file,rows=self.book();self.assertEqual(rows[1][23],'Reviewed Contributor')
  wb=openpyxl.load_workbook(file);self.assertEqual(wb['SP 2'].tables['TitleVisionErrors'].ref,'A1:AC4')
  self.assertTrue(wb['SP 2'].column_dimensions['AC'].hidden);self.assertEqual(len(wb['Summary']._pivots),4)
  self.assertEqual(wb['Summary']._pivots[0].cache.recordCount,3)
  self.assertIn('Reviewed Contributor',wb['SP 2']['X3'].value)
  fixture=ROOT/'tests/results'/('report-edits-'+uuid.uuid4().hex);fixture.mkdir();shutil.copy2(file,fixture/'report.xlsx')
  print('EDITS_FIXTURE='+str(fixture/'report.xlsx'))

 def test_old_copy_does_not_overwrite_newer_correction_or_delete_new_rows(self):
  s=self.source(1,'2026-10-01');self.live=[s];self.add(self.seed('2026-10-01','2026-10-01',[s]),self.real_build)
  old=self.edited_copy(lambda wb:setattr(wb['SP 2']['S2'],'value','Old correction'))
  new=self.edited_copy(lambda wb:setattr(wb['SP 2']['S2'],'value','New correction'))
  manual_edits.save(ROOT,self.data,self.db,self.real_build,'2026-10',new)
  before=self.book()[0]
  with self.assertRaisesRegex(ValueError,'conflicts'):manual_edits.save(ROOT,self.data,self.db,self.real_build,'2026-10',old)
  self.assertEqual(self.book()[0],before)
  self.assertEqual(self.book()[1][0][18],'New correction')

 def test_invalid_upload_and_build_failure_leave_saved_corrections_unchanged(self):
  s=self.source(1,'2026-10-01');self.live=[s];self.add(self.seed('2026-10-01','2026-10-01',[s]),self.real_build)
  cases=[lambda wb:setattr(wb['SP 2']['AC2'],'value','999'),lambda wb:setattr(wb['SP 2']['B2'],'value','Other order'),
   lambda wb:setattr(wb['SP 2']['W2'],'value','Other team'),lambda wb:setattr(wb['SP 2']['S2'],'value','=1+1'),lambda wb:wb['SP 2'].delete_rows(2)]
  before=self.book()[0]
  for change in cases:
   with self.subTest(change=change):
    with self.assertRaises(ValueError):manual_edits.save(ROOT,self.data,self.db,self.real_build,'2026-10',self.edited_copy(change))
    self.assertEqual(self.book()[0],before);self.assertEqual(manual_edits.load(self.db,'2026-10'),{})
  upload=self.edited_copy(lambda wb:setattr(wb['SP 2']['S2'],'value','Correction'))
  def fail(args,folder):raise ValueError('Build failed')
  with self.assertRaisesRegex(ValueError,'Build failed'):manual_edits.save(ROOT,self.data,self.db,fail,'2026-10',upload)
  self.assertEqual(self.book()[0],before);self.assertEqual(manual_edits.load(self.db,'2026-10'),{})

 def test_unchanged_and_stale_unedited_copy_keeps_manual_changes(self):
  s=self.source(1,'2026-10-01');self.live=[s];self.add(self.seed('2026-10-01','2026-10-01',[s]),self.real_build)
  untouched=self.book()[0].read_bytes()
  upload=self.edited_copy(lambda wb:setattr(wb['SP 2']['X2'],'value','Manual final'))
  manual_edits.save(ROOT,self.data,self.db,self.real_build,'2026-10',upload)
  # Opening an old workbook without edits cannot revert a saved correction.
  result=manual_edits.save(ROOT,self.data,self.db,self.real_build,'2026-10',untouched)
  self.assertIn('No new',result['message']);self.assertEqual(self.book()[1][0][23],'Manual final')
  # A changed Team uses the other team's contributor instead of the old override.
  upload=self.edited_copy(lambda wb:setattr(wb['SP 2']['W2'],'value','Type'))
  manual_edits.save(ROOT,self.data,self.db,self.real_build,'2026-10',upload)
  self.assertEqual(self.book()[1][0][23],'Typer')

 def test_older_copy_can_save_nonconflicting_edits_without_dropping_next_day(self):
  first=self.source(1,'2026-10-01');self.live=[first]
  self.add(self.seed('2026-10-01','2026-10-01',[first]),self.real_build)
  upload=self.edited_copy(lambda wb:setattr(wb['SP 2']['S2'],'value','Reviewed original'))
  second=self.source(2,'2026-10-02');self.live.append(second)
  self.add(self.seed('2026-10-02','2026-10-02',[second]),self.real_build)
  manual_edits.save(ROOT,self.data,self.db,self.real_build,'2026-10',upload)
  rows=self.book()[1];self.assertEqual(len(rows),2)
  self.assertEqual([r[23] for r in rows],['Reviewed original','Searcher'])
  self.assertEqual(self.book()[0].parent.joinpath('payload.json').exists(),True)

if __name__=='__main__':unittest.main()
