"""Regressions for native PivotTable packaging and legacy report downloads."""
import json,sys,subprocess,uuid,unittest,zipfile
from pathlib import Path
from lxml import etree as E
app=Path(__file__).resolve().parents[1]
S='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
ns={'x':S}
class ReportTests(unittest.TestCase):
 def test_historical_team_edits_keep_their_final_name(self):
  sys.path.insert(0,str(app))
  import report_layout,manual_edits
  row=['']*27;row[3]='Type';row[18]='Searcher';row[20]='Typer';row[22]='Typer'
  migrated=report_layout.current_rows(report_layout.PREVIOUS_HEADERS,[row])
  self.assertEqual(migrated[0][3],'Typing')
  rows,overrides=manual_edits.apply(migrated,[{'sourceId':'1'}],{'1':{'3':'Type','final':{'team':'Type','name':'Reviewed name'}}})
  self.assertEqual((rows[0][3],rows[0][22]),('Typing','Reviewed name'))
  self.assertEqual(overrides[0]['team'],'Typing')
  for team,want in [('Search','Searcher'),('Typing','Typer'),('Triage','Triage'),('VM Team','VM Team'),('Leadership','Leadership')]:
   rows,_=manual_edits.apply(migrated,[{'sourceId':'1'}],{'1':{'3':team}})
   self.assertEqual(rows[0][22],want)
 def test_report_scenarios(self):
  layout=json.loads((app/'templates/layout.json').read_text(encoding='utf8'))
  for scenario,teams in [('empty',[]),('single',['Search']),('all-teams',['Search','Typing','Triage','VM Team','Leadership',''])]:
   with self.subTest(scenario=scenario):
    folder=app/'tests/results'/('report-'+scenario+'-'+uuid.uuid4().hex);folder.mkdir(parents=True)
    rows=[]
    for i,team in enumerate(teams):
     r=['']*27;r[0]='Non-Chargeable';r[1]='Test-'+str(i);r[6]='Test category';r[8]='=Literal source text';r[10]=r[14]=r[15]=46282;r[12]='CA';r[13]=i;r[18]='Example Searcher';r[20]='Example Typer';r[3]=team;r[22]=r[18] if team=='Search' else r[20] if team=='Typing' else team;r[23:27]=['Disputed','=Literal vendor comment\nSecond line','Non-Chargeable','Training, reason accepted'];rows.append(r)
    specs=[]
    for name,field,title,col,row in [('ErrorsByContributor',22,'Final Error Contributor',0,5),('ErrorsByTeam',3,'Team',4,5),('ErrorsByCategory',6,'Error Category',8,5),('ErrorsByStatus',0,'Status',4,15)]:
     items=sorted({r[field] for r in rows}|({'VM team'} if field in (3,22) else set()))
     specs.append(dict(name=name,field=field,title=title,col=col,row=row,items=items,height=len(items)+3))
    # Deliberately omit widths to exercise rebuilding older saved payloads.
    (folder/'payload.json').write_text(json.dumps({'data':[layout['headers'],*rows],'specs':specs,'tableName':'TitleVisionErrors'}),encoding='utf8')
    (folder/'collection.json').write_text(json.dumps({'start':'2026-09-17','end':'2026-09-17','count':len(rows),'points':sum(r[13] for r in rows)}),encoding='utf8')
    for script in ('build_portable.py','finish.py'):
     result=subprocess.run([sys.executable,str(app/'report'/script),str(folder)],capture_output=True,text=True)
     self.assertEqual(result.returncode,0,result.stderr)
    with zipfile.ZipFile(folder/'report.xlsx') as z:
     sheet=E.fromstring(z.read('xl/worksheets/sheet1.xml'));self.assertIsNone(sheet.find('x:pivotParts',ns))
     rels=E.fromstring(z.read('xl/worksheets/_rels/sheet1.xml.rels'));self.assertEqual(sum(r.get('Type','').endswith('/pivotTable') for r in rels),4)
     styles=E.fromstring(z.read('xl/styles.xml'))
     for style in styles.findall('x:tableStyles/x:tableStyle',ns):self.assertFalse(any(k.startswith('{') for k in style.attrib))
     order=[E.QName(c).localname for c in E.fromstring(z.read('xl/workbook.xml'))];self.assertLess(order.index('calcPr'),order.index('pivotCaches'))
     cache=E.fromstring(z.read('xl/pivotCache/pivotCacheDefinition1.xml'))
     fields=cache.find('x:cacheFields',ns)
     records=E.fromstring(z.read('xl/pivotCache/pivotCacheRecords1.xml'))
     decoded=[]
     for record in records:
      decoded.append([fields[j].find('x:sharedItems',ns)[int(c.get('v'))].get('v') if E.QName(c).localname=='x' else c.get('v') for j,c in enumerate(record)])
     self.assertEqual([float(r[13]) for r in decoded],[r[13] for r in rows])
     self.assertEqual([r[1] for r in decoded],[r[1] for r in rows])
     self.assertEqual([r[22] for r in decoded],['Non-Chargeable']*len(rows))
     pivots=[E.fromstring(z.read('xl/pivotTables/pivotTable'+str(k)+'.xml')) for k in range(1,5)]
     self.assertEqual([p.get('rowHeaderCaption') for p in pivots],['Names','Status','Team','Points'])
     for pivot in pivots:
      self.assertEqual(pivot.get('enableDrill'),'1')
      self.assertEqual(pivot.find('x:dataFields',ns)[0].get('fld'),'12')
     for shared in cache.findall('x:cacheFields/x:cacheField/x:sharedItems',ns):
      if shared.get('count') is not None:self.assertEqual(shared.get('containsBlank')=='1',shared.find('x:m',ns) is not None)
    self.assertEqual(json.loads((folder/'validation.json').read_text())['workbookFormatVersion'],4)
    if rows:
     import openpyxl
     sheet=openpyxl.load_workbook(folder/'report.xlsx')['SP 2']
     self.assertEqual([sheet.cell(2,c).value for c in range(24,28)],rows[0][23:27]);self.assertEqual(sheet['Y2'].data_type,'s');self.assertFalse(sheet['AA2'].alignment.wrap_text)
     self.assertEqual(sheet['D1'].value,'Team');self.assertEqual(sheet['D2'].value,teams[0]);self.assertEqual(sheet['W2'].data_type,'f')
     self.assertEqual(sheet['W1'].value,'Final Error Contributor')
     self.assertTrue(sheet['W2'].value.startswith('=IF(TRIM(A2)="Non-Chargeable","Non-Chargeable",'))
     cached=openpyxl.load_workbook(folder/'report.xlsx',data_only=True)
     self.assertEqual(cached['SP 2']['W2'].value,'Non-Chargeable')
     self.assertEqual(cached['Sheet1']['A4'].value,'Non-Chargeable')
     self.assertEqual(cached['Sheet1']['B4'].value,len(rows))
     self.assertEqual(cached['Sheet1']['C4'].value,sum(r[13] for r in rows))
     cached.close()
     self.assertEqual(sheet.data_validations.dataValidation[0].formula1,'"Search,Typing,Triage,VM Team,Leadership"')
     self.assertEqual(sheet.freeze_panes,'A2')
     self.assertEqual(sheet.row_dimensions[1].height,18.75)
    print('REPORT_FIXTURE='+str(folder/'report.xlsx'))
if __name__=='__main__':unittest.main()
